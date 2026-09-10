"""
Generación del Excel institucional.

Las etiquetas y el orden ya NO se definen aquí. Se derivan de
app.domain.fields, la fuente de verdad única. Antes este módulo mantenía su
propia copia de los 25 campos y ya había divergido de FIELD_HEADERS en dos
etiquetas ("Duración Total de la Sesión" y "% de Cumplimiento de Evaluación").
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.fields import FIELDS as _CONTRACT_FIELDS

# Colores institucionales EPM (hex sin #)
_GREEN = "00A650"
_BLUE = "0066B3"
_DARK_BLUE = "003B71"
_WHITE = "FFFFFF"
_GREEN_LIGHT = "E8F5E9"
_BLUE_LIGHT = "E3F2FD"
_PURPLE_LIGHT = "EDE7F6"

# (etiqueta, clave) por bloque, derivado del contrato. Imposible desincronizar.
_FIELDS = {
    f"block{b}": [(f.header, f.key) for f in _CONTRACT_FIELDS if f.block == b]
    for b in (1, 2, 3)
}


def _section_header(ws, row: int, text: str, color: str, cols: int = 16) -> None:
    end_col = get_column_letter(cols)
    ws.merge_cells(f"A{row}:{end_col}{row}")
    cell = ws[f"A{row}"]
    cell.value = text
    cell.font = Font(bold=True, color=_WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 22


def _data_row(ws, row: int, label: str, value: str, bg: str, cols: int = 16) -> None:
    end_col = get_column_letter(cols)
    label_cell = ws[f"A{row}"]
    label_cell.value = label
    label_cell.font = Font(bold=True, size=10)
    label_cell.fill = PatternFill("solid", fgColor=bg)
    label_cell.alignment = Alignment(vertical="center", indent=1, wrap_text=True)

    ws.merge_cells(f"B{row}:{end_col}{row}")
    value_cell = ws[f"B{row}"]
    value_cell.value = value
    value_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = max(20, min(60, len(value) // 4 + 20) if value else 20)


def generate_excel(form_data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO"

    # === Title row ===
    ws.merge_cells("A1:P1")
    title = ws["A1"]
    title.value = "FUNDACIÓN EPM — CONSOLIDACIÓN METODOLÓGICA"
    title.font = Font(bold=True, color=_WHITE, size=14)
    title.fill = PatternFill("solid", fgColor=_DARK_BLUE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:P2")
    sub = ws["A2"]
    sub.value = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Fundación EPM"
    sub.font = Font(italic=True, color="888888", size=9)
    sub.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.row_dimensions[2].height = 16

    row = 3

    # === Block 1 ===
    _section_header(ws, row, "BLOQUE 1 — IDENTIFICACIÓN Y DISEÑO METODOLÓGICO", _GREEN)
    row += 1
    for label, key in _FIELDS["block1"]:
        _data_row(ws, row, label, form_data.get(key, ""), _GREEN_LIGHT)
        row += 1

    # === Block 2 ===
    _section_header(ws, row, "BLOQUE 2 — INFORME DE EJECUCIÓN", _BLUE)
    row += 1
    for label, key in _FIELDS["block2"]:
        _data_row(ws, row, label, form_data.get(key, ""), _BLUE_LIGHT)
        row += 1

    # === Block 3 ===
    _section_header(ws, row, "BLOQUE 3 — EVALUACIÓN", _DARK_BLUE)
    row += 1
    for label, key in _FIELDS["block3"]:
        _data_row(ws, row, label, form_data.get(key, ""), _PURPLE_LIGHT)
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 34
    for col in range(2, 17):
        ws.column_dimensions[get_column_letter(col)].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

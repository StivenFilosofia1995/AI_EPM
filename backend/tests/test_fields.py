"""
Sincronía del contrato de 25 campos.

Estas pruebas existen porque el diseño anterior tenía TRES copias de la lista
de campos (google_sheets_service, chat_service y excel_service) y ya habían
divergido en dos etiquetas. Si alguien vuelve a duplicarlas, esto falla.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from app.domain import fields as F

SQL_001 = Path(__file__).resolve().parent.parent / "sql" / "migrations" / "001_esquema_base.sql"


def test_son_exactamente_25_campos():
    assert len(F.FIELDS) == 25
    assert len(F.FIELD_KEYS) == 25
    assert len(F.FIELD_HEADERS) == 25


def test_claves_y_encabezados_sin_duplicados():
    assert len(set(F.FIELD_KEYS)) == 25, "Hay claves de campo repetidas."
    assert len(set(F.FIELD_HEADERS)) == 25, "Hay encabezados repetidos."


def test_distribucion_por_bloques():
    assert F.BLOCK_FIELD_COUNTS == {1: 16, 2: 4, 3: 5}


def test_assert_contract_no_lanza():
    F.assert_contract()


def test_claves_y_encabezados_alineados_uno_a_uno():
    for i, campo in enumerate(F.FIELDS):
        assert F.FIELD_KEYS[i] == campo.key
        assert F.FIELD_HEADERS[i] == campo.header


def test_excel_service_deriva_del_contrato():
    """excel_service ya no puede tener su propia lista."""
    from app.services import excel_service as ex

    todas = [par for bloque in ("block1", "block2", "block3") for par in ex._FIELDS[bloque]]
    assert [k for _, k in todas] == F.FIELD_KEYS
    assert [h for h, _ in todas] == F.FIELD_HEADERS


def test_ya_no_existe_el_servicio_de_google_sheets():
    """La integración con Sheets se retiró: la exportación es Excel."""
    import importlib

    try:
        importlib.import_module("app.services.google_sheets_service")
    except ModuleNotFoundError:
        return
    raise AssertionError("google_sheets_service debería haberse eliminado.")


def test_el_excel_por_lote_respeta_el_orden_del_contrato():
    """Las 25 columnas van primero y en orden; el seguimiento va después."""
    import io

    from openpyxl import load_workbook

    from app.services.excel_service import generate_excel_lote

    libro = load_workbook(io.BytesIO(generate_excel_lote([])))
    hoja = libro.active
    cabeceras = [hoja.cell(1, i).value for i in range(1, len(F.FIELD_HEADERS) + 1)]
    assert cabeceras == F.FIELD_HEADERS
    assert hoja.cell(1, len(F.FIELD_HEADERS) + 1).value == "Facilitador"


def test_email_service_usa_el_mismo_orden():
    from app.services import email_service as em

    assert em.FIELD_KEYS == F.FIELD_KEYS
    assert em._KEY_TO_HEADER == dict(zip(F.FIELD_KEYS, F.FIELD_HEADERS, strict=True))


def test_columnas_de_la_migracion_coinciden_con_las_claves():
    """
    El orden de las 25 columnas en epm_actividades debe ser el del contrato.
    Es la exportación institucional la que depende de ese orden.
    """
    sql = SQL_001.read_text(encoding="utf-8")
    bloque = sql[sql.index("CREATE TABLE IF NOT EXISTS public.epm_actividades") :]
    bloque = bloque[: bloque.index(");")]

    encontradas = []
    for linea in bloque.splitlines():
        limpia = linea.strip()
        if not limpia or limpia.startswith("--"):
            continue
        m = re.match(r"^([a-z_]+)\s+(TEXT|DATE|INTEGER|SMALLINT|UUID|TIMESTAMPTZ)", limpia)
        if m and m.group(1) in F.FIELD_KEYS:
            encontradas.append(m.group(1))

    assert encontradas == F.FIELD_KEYS, (
        "El orden de las columnas de epm_actividades no coincide con FIELD_KEYS."
    )


def test_hay_17_ods():
    assert len(F.ODS) == 17
    assert F.ODS[0].startswith("1.")
    assert F.ODS[-1].startswith("17.")


@pytest.mark.parametrize(
    "catalogo,esperados",
    [("programa", 4), ("linea_accion", 4), ("tipo_actividad", 6), ("publico", 6), ("ods", 17)],
)
def test_catalogos_cerrados(catalogo, esperados):
    assert len(F.CLOSED_OPTIONS[catalogo]) == esperados


def test_serializacion_de_seleccion_multiple():
    valores = [F.ODS[3], F.ODS[5]]
    texto = F.serialize_multi(valores)
    assert texto == f"{F.ODS[3]}; {F.ODS[5]}"
    assert F.deserialize_multi(texto) == valores


def test_lugar_no_tiene_catalogo_todavia():
    """Guarda contra inventar un catálogo institucional que no fue entregado."""
    assert "lugar" not in F.CLOSED_OPTIONS
